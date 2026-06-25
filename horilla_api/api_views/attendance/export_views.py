"""Export attendance activity data as JSON preview or Excel (.xlsx)."""
import io
from datetime import date, timedelta

from django.db.models import Min, Max
from django.http import HttpResponse
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from attendance.models import Attendance, AttendanceActivity
from employee.models import Employee


DAY_NAMES_VI = {
    0: "T2", 1: "T3", 2: "T4", 3: "T5", 4: "T6", 5: "T7", 6: "CN",
}


def _is_cnb(user):
    if user.is_superuser:
        return True
    gnames = [g.name.lower() for g in user.groups.all()]
    return any(
        "c&b" in g or "c & b" in g or "chuyên viên c" in g or "cb" == g.strip()
        for g in gnames
    )


def _parse_d(s):
    """Parse 'YYYY-MM-DD' → date, hoặc None nếu rỗng/sai định dạng."""
    if not s:
        return None
    try:
        return date.fromisoformat(s.strip())
    except (ValueError, AttributeError):
        return None


def _build_rows(year, month, company_id=None, department_id=None, search=None,
                from_date=None, to_date=None):
    import calendar
    from django.db.models import Q, F, Value, CharField
    from django.db.models.functions import Concat

    # Khoảng ngày: mặc định cả tháng; nếu có from/to thì kẹp trong phạm vi tháng.
    month_first = date(year, month, 1)
    month_last = date(year, month, calendar.monthrange(year, month)[1])
    first = from_date or month_first
    last = to_date or month_last
    if first < month_first:
        first = month_first
    if last > month_last:
        last = month_last

    employees = Employee.objects.filter(is_active=True).select_related(
        "employee_work_info__shift_id"
    ).order_by("stt", "employee_first_name")
    if company_id:
        employees = employees.filter(employee_work_info__company_id=company_id)
    if department_id:
        employees = employees.filter(employee_work_info__department_id=department_id)

    # Lọc theo Tên / Họ đệm+Tên / Mã NV (badge) / Mã Kế toán. Nhiều NV: phân cách dấu phẩy.
    if search and search.strip():
        terms = [t.strip() for t in search.split(",") if t.strip()] or [search.strip()]
        employees = employees.annotate(
            _fullname=Concat(
                F("employee_last_name"), Value(" "), F("employee_first_name"),
                output_field=CharField(),
            )
        )
        cond = Q()
        for t in terms:
            cond |= (
                Q(employee_first_name__icontains=t) | Q(employee_last_name__icontains=t)
                | Q(_fullname__icontains=t) | Q(badge_id__icontains=t)
                | Q(accounting_code__icontains=t)
            )
        employees = employees.filter(cond)

    rows = []
    stt = 0
    for emp in employees:
        wi = getattr(emp, "employee_work_info", None)
        shift = wi.shift_id if wi else None

        atts = Attendance.objects.filter(
            employee_id=emp,
            attendance_date__gte=first,
            attendance_date__lte=last,
        ).order_by("attendance_date")

        for att in atts:
            stt += 1
            d = att.attendance_date
            weekday = DAY_NAMES_VI.get(d.weekday(), "")

            acts = AttendanceActivity.objects.filter(
                employee_id=emp,
                attendance_date=d,
            ).order_by("clock_in")

            # ALD26: gom mọi lượt chấm phẳng (clock_in + clock_out), sắp theo thời gian.
            # Lượt 1 = giờ vào; lượt cuối (ngày đã qua) = giờ ra; chỉ 1 lượt → NCO.
            punch_times = []
            for act in acts:
                if act.clock_in:
                    punch_times.append(act.clock_in)
                if act.clock_out:
                    punch_times.append(act.clock_out)
            punch_times.sort()

            def _hm(t):
                s = str(t)[:5]
                return "NCO" if s == "23:59" else s

            detail = " · ".join(_hm(t) for t in punch_times)

            earliest_in = punch_times[0] if punch_times else None
            latest_out = punch_times[-1] if len(punch_times) >= 2 else None

            worked = att.attendance_worked_hour or "00:00"
            earliest_str = str(earliest_in)[:5] if earliest_in else "--:--"
            if len(punch_times) >= 2:
                latest_str = _hm(punch_times[-1])
            elif len(punch_times) == 1 and d < date.today():
                latest_str = "NCO"
            else:
                latest_str = "--:--"

            is_late = False
            is_early = False
            late_mins = 0
            early_mins = 0
            coefficient = 1.0
            min_hour_str = "08:00"
            notes = []

            if shift:
                from base.models import EmployeeShiftSchedule, EmployeeShiftDay
                day_name = d.strftime("%A").lower()
                try:
                    day_obj = EmployeeShiftDay.objects.get(day=day_name)
                    sched = EmployeeShiftSchedule.objects.filter(
                        shift_id=shift, day=day_obj
                    ).first()
                    if sched:
                        coefficient = float(sched.work_day_coefficient)
                        min_hour_str = sched.minimum_working_hour or "08:00"

                        grace_secs = 0
                        if shift.grace_time_id:
                            grace_secs = shift.grace_time_id.allowed_time_in_secs or 0

                        # Giờ chuẩn tính trễ/sớm: ưu tiên giờ lõi (core_*) nếu được
                        # khai báo (vd ALD26: 08:00/17:30), nếu không dùng khung ca.
                        late_ref = sched.core_start_time or sched.start_time
                        early_ref = sched.core_end_time or sched.end_time

                        if late_ref and earliest_in:
                            from datetime import datetime as _dt, timedelta as _td
                            shift_start_dt = _dt.combine(d, late_ref) + _td(seconds=grace_secs)
                            check_in_dt = _dt.combine(d, earliest_in)
                            if check_in_dt > shift_start_dt:
                                is_late = True
                                late_mins = int((check_in_dt - shift_start_dt).total_seconds() / 60)
                                notes.append(f"Trễ {late_mins} phút")

                        if early_ref and latest_out:
                            from datetime import datetime as _dt
                            shift_end_dt = _dt.combine(d, early_ref)
                            check_out_dt = _dt.combine(d, latest_out)
                            if check_out_dt < shift_end_dt:
                                is_early = True
                                early_mins = int((shift_end_dt - check_out_dt).total_seconds() / 60)
                                notes.append(f"Sớm {early_mins} phút")
                except Exception:
                    pass

            worked_secs = 0
            try:
                parts = worked.split(":")
                worked_secs = int(parts[0]) * 3600 + int(parts[1]) * 60
                if len(parts) > 2:
                    worked_secs += int(parts[2])
            except Exception:
                pass

            min_parts = min_hour_str.split(":")
            min_secs = int(min_parts[0]) * 3600 + int(min_parts[1]) * 60
            standard_secs = int(min_secs * coefficient) if coefficient else min_secs

            pct = round((worked_secs / standard_secs * 100), 1) if standard_secs > 0 else 0
            cong = round(min(1.0, worked_secs / standard_secs), 2) if standard_secs > 0 else 0.0
            if pct < 100 and worked_secs > 0:
                notes.append(f"Đạt {pct}% ngày công")
            elif pct > 100:
                notes.append(f"Đạt {pct}% ngày công")

            note_str = " | ".join(notes)

            rows.append({
                "stt": stt,
                "employee_code": emp.employee_code or emp.badge_id or "",
                "accounting_code": emp.accounting_code or "",
                "first_name": emp.employee_first_name,
                "full_name": f"{emp.employee_last_name or ''} {emp.employee_first_name}".strip(),
                "date": d.isoformat(),
                "weekday": weekday,
                "clock_in": earliest_str,
                "clock_out": latest_str,
                "worked": worked[:5] if len(worked) >= 5 else worked,
                "detail": detail,
                "is_late": is_late,
                "is_early": is_early,
                "late_mins": late_mins,
                "early_mins": early_mins,
                "coefficient": coefficient,
                "work_pct": pct,
                "cong": cong,
                "note": note_str,
            })

    return rows


class AttendanceExportPreviewView(APIView):
    """GET: Preview attendance data for a month (JSON)."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not _is_cnb(request.user):
            return Response({"error": "Chỉ C&B mới được xuất"}, status=403)

        year = int(request.query_params.get("year", date.today().year))
        month = int(request.query_params.get("month", date.today().month))
        company_id = request.query_params.get("company_id") or None
        department_id = request.query_params.get("department_id") or None
        search = request.query_params.get("q") or None

        import calendar
        month_first = date(year, month, 1)
        month_last = date(year, month, calendar.monthrange(year, month)[1])
        from_date = _parse_d(request.query_params.get("from_date")) or month_first
        to_date = _parse_d(request.query_params.get("to_date")) or month_last
        from_date = max(from_date, month_first)
        to_date = min(to_date, month_last)

        try:
            page = max(1, int(request.query_params.get("page", 1) or 1))
        except (TypeError, ValueError):
            page = 1
        try:
            page_size = int(request.query_params.get("page_size", 50) or 50)
        except (TypeError, ValueError):
            page_size = 50
        if page_size not in (20, 50, 100, 200):
            page_size = 50

        rows = _build_rows(year, month, company_id, department_id, search, from_date, to_date)
        total = len(rows)
        total_pages = max(1, (total + page_size - 1) // page_size)
        page = min(page, total_pages)
        start = (page - 1) * page_size
        page_rows = rows[start:start + page_size]

        return Response({
            "year": year,
            "month": month,
            "from_date": from_date.isoformat(),
            "to_date": to_date.isoformat(),
            "month_first": month_first.isoformat(),
            "month_last": month_last.isoformat(),
            "count": total,
            "page": page,
            "page_size": page_size,
            "total_pages": total_pages,
            "results": page_rows,
        })


class AttendanceExportExcelView(APIView):
    """GET: Download attendance data as .xlsx file."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not _is_cnb(request.user):
            return Response({"error": "Chỉ C&B mới được xuất"}, status=403)

        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        year = int(request.query_params.get("year", date.today().year))
        month = int(request.query_params.get("month", date.today().month))
        company_id = request.query_params.get("company_id") or None
        department_id = request.query_params.get("department_id") or None
        search = request.query_params.get("q") or None

        import calendar
        month_first = date(year, month, 1)
        month_last = date(year, month, calendar.monthrange(year, month)[1])
        from_date = _parse_d(request.query_params.get("from_date")) or month_first
        to_date = _parse_d(request.query_params.get("to_date")) or month_last
        from_date = max(from_date, month_first)
        to_date = min(to_date, month_last)

        rows = _build_rows(year, month, company_id, department_id, search, from_date, to_date)

        wb = Workbook()
        ws = wb.active
        ws.title = f"CC T{month}-{year}"

        headers = [
            "STT", "Mã N.Viên", "Mã KT", "Tên", "Họ tên đầy đủ",
            "Ngày", "Thứ", "Giờ vào", "Giờ ra", "Giờ làm",
            "Lượt chấm", "Đi trễ", "Về sớm",
            "Hệ số", "% Ngày công", "Công", "Ghi chú",
        ]

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="142B6F", end_color="142B6F", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        late_fill = PatternFill(start_color="FDE2E2", end_color="FDE2E2", fill_type="solid")
        early_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")

        for i, r in enumerate(rows, 2):
            vals = [
                r["stt"], r["employee_code"], r["accounting_code"],
                r["first_name"], r["full_name"],
                r["date"], r["weekday"], r["clock_in"], r["clock_out"],
                r["worked"], r["detail"],
                "Có" if r["is_late"] else "",
                "Có" if r["is_early"] else "",
                r["coefficient"], f'{r["work_pct"]}%', r["cong"],
                r["note"],
            ]
            for col, v in enumerate(vals, 1):
                cell = ws.cell(row=i, column=col, value=v)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")
                if r["is_late"] and col == 12:
                    cell.fill = late_fill
                if r["is_early"] and col == 13:
                    cell.fill = early_fill

        col_widths = [6, 14, 10, 15, 25, 12, 6, 10, 10, 10, 35, 8, 8, 8, 12, 8, 30]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[chr(64 + i) if i <= 26 else ""].width = w

        ws.auto_filter.ref = ws.dimensions

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        filename = f"HoatDong_ChamCong_T{month}_{year}.xlsx"
        response = HttpResponse(
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response
