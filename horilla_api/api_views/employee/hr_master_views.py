"""App Feature 'Data Nhân sự' + 'Danh mục' — HR Master Data.

Bảng nhân sự đầy đủ theo mẫu Master Data (69 cột / 9 nhóm). Field map thẳng vào
Employee / EmployeeWorkInformation / HNHEmployeeProfile / EmployeeBankDetails;
các cột chưa có model (hợp đồng 3 lần, đánh giá, offboarding, tuyển dụng...) lưu
trong Employee.additional_info['hr_master'] (JSON) — vẫn hiển thị + sửa được.

Endpoints (đều yêu cầu quyền HR/C&B):
  GET  /api/employee/hr-master/            list + filter (q, company_id, department_id, join_year), phân trang
  GET  /api/employee/hr-master/<pk>/       chi tiết đầy đủ 1 NV
  PUT  /api/employee/hr-master/<pk>/       cập nhật (atomic, 4 model + extras)
  GET  /api/employee/hr-master/export/     xuất Excel lưới 69 cột (theo filter)
  GET  /api/employee/hr-categories/        Danh mục: công ty, phòng ban, vị trí + enum
"""
from django.db import transaction
from django.db.models import Q
from rest_framework.pagination import PageNumberPagination
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from employee.models import (
    Employee,
    EmployeeWorkInformation,
    EmployeeBankDetails,
    HNHEmployeeProfile,
)

# ── Danh mục enum (nguồn: sheet 'Danh mục' của Master Data) ──────────────────
ENUM_CATEGORIES = {
    "gender": ["Nam", "Nữ"],
    "marital_status": ["Độc thân", "Đã kết hôn", "Ly hôn"],
    "recruit_source": ["Giới thiệu nội bộ", "Website tuyển dụng", "LinkedIn",
                       "Headhunter", "Trường/Job fair", "Facebook", "Khác"],
    "probation_status": ["Đang thử việc", "Đạt - Ký chính thức", "Không đạt - Chấm dứt"],
    "contract_type": ["HĐ thử việc", "HĐ xác định thời hạn (dưới 12 tháng)",
                      "HĐ xác định thời hạn (12-36 tháng)", "HĐ không xác định thời hạn",
                      "HĐ thời vụ/Cộng tác viên"],
    "work_form": ["Full-time", "Part-time", "Theo ca", "Từ xa (Remote)"],
    "pay_method": ["Chuyển khoản", "Tiền mặt"],
    "education": ["THPT", "Trung cấp", "Cao đẳng", "Đại học", "Sau đại học (Thạc sĩ/Tiến sĩ)"],
    "eval_rating": ["Xuất sắc (A+)", "Tốt (A)", "Đạt yêu cầu (B)", "Cần cải thiện (C)", "Không đạt (D)"],
    "work_status": ["Đang làm việc", "Đã nghỉ việc", "Nghỉ thai sản",
                    "Nghỉ không lương/Tạm hoãn HĐ"],
    "resign_type": ["Tự nguyện xin nghỉ", "Hết hạn hợp đồng - không tái ký",
                    "Công ty chấm dứt HĐ/Sa thải", "Công ty tái cấu trúc", "Nghỉ hưu"],
    "handover": ["Chưa bắt đầu", "Đang bàn giao", "Đã hoàn tất"],
    "nationality": ["Việt Nam"],
}
ENUM_LABELS = {
    "gender": "Giới tính", "marital_status": "Tình trạng hôn nhân",
    "recruit_source": "Nguồn tuyển dụng", "probation_status": "Trạng thái thử việc",
    "contract_type": "Loại hợp đồng", "work_form": "Hình thức làm việc",
    "pay_method": "Hình thức trả lương", "education": "Trình độ học vấn",
    "eval_rating": "Xếp loại đánh giá", "work_status": "Trạng thái làm việc",
    "resign_type": "Loại nghỉ việc", "handover": "Bàn giao công việc",
    "nationality": "Quốc tịch",
}

# Các field lưu trong additional_info['hr_master'] (không có cột model riêng).
EXTRA_FIELDS = [
    "recruit_applied_date", "recruit_source", "probation_end", "probation_salary",
    "probation_status", "contract1_type", "contract1_no", "contract1_start", "contract1_end",
    "contract2_type", "contract2_no", "contract2_start", "contract2_end",
    "contract3_type", "contract3_no", "contract3_start", "work_form",
    "allowance", "pay_method", "dependents_names", "school",
    "eval_rating", "raise_date", "raise_amount", "raise_effective", "promotion_date",
    "work_status", "resign_date", "resign_reason", "resign_type", "notice_date",
    "handover", "note", "team", "level_label", "nationality", "company_code", "dept_code",
    # Ngày hệ thống ghi khi C&B chuyển Tạm nghỉ (auto). Dùng cho export NV nghỉ/tháng.
    "deactivated_date",
]

# Thời gian thử việc (ngày) tính từ ngày vào làm. Đánh dấu "Đang thử việc" theo
# THÔNG TIN NHÂN VIÊN (date_joining), KHÔNG theo loại hợp đồng (thường bị cũ).
PROBATION_DAYS = 60  # 2 tháng (chuẩn Luật LĐ VN cho đa số vị trí)


def _can_view(request) -> bool:
    u = request.user
    if u.is_superuser or u.has_perm("employee.view_employee"):
        return True
    for g in u.groups.all():
        n = (g.name or "").lower()
        if "c&b" in n or "chuyên viên c" in n or "nhân sự" in n:
            return True
    return False


def _can_edit(request) -> bool:
    u = request.user
    if u.is_superuser or u.has_perm("employee.change_employee"):
        return True
    for g in u.groups.all():
        n = (g.name or "").lower()
        if "c&b" in n or "chuyên viên c" in n or "nhân sự" in n:
            return True
    return False


def _vn_name(emp) -> str:
    return f"{emp.employee_first_name} {emp.employee_last_name or ''}".strip()


def _d(v):
    return v.isoformat() if v else None


def _extra(emp) -> dict:
    ai = emp.additional_info if isinstance(emp.additional_info, dict) else {}
    hm = ai.get("hr_master") if isinstance(ai.get("hr_master"), dict) else {}
    return hm


def serialize_employee(emp) -> dict:
    """Dict đầy đủ 69 cột cho 1 NV (map model + extras)."""
    wi = getattr(emp, "employee_work_info", None)
    prof = getattr(emp, "hnh_profile", None)
    bank = getattr(emp, "employee_bank_details", None)
    hm = _extra(emp)
    dept = wi.department_id if wi else None
    comp = wi.company_id if wi else None
    pos = wi.job_position_id if wi else None
    rm = wi.reporting_manager_id if wi else None
    return {
        "id": emp.id,
        # Cơ cấu tổ chức
        "company_code": hm.get("company_code") or "",
        "company": comp.company if comp else "",
        "company_id": comp.id if comp else None,
        "dept_code": hm.get("dept_code") or "",
        "department": dept.department if dept else "",
        "department_id": dept.id if dept else None,
        "team": hm.get("team") or "",
        # Thông tin nhân viên
        "badge_id": emp.badge_id or emp.employee_code or "",
        "accounting_code": emp.accounting_code or "",
        "name": _vn_name(emp),
        "first_name": emp.employee_first_name or "",
        "last_name": emp.employee_last_name or "",
        "dob": _d(emp.dob),
        "gender": emp.gender or "",
        "cccd": (prof.cccd if prof else "") or "",
        "cccd_issue_date": _d(prof.cccd_issue_date) if prof else None,
        "cccd_issue_place": (prof.cccd_issue_place if prof else "") or "",
        "nationality": hm.get("nationality") or emp.country or "",
        "marital_status": emp.marital_status or "",
        # Liên hệ
        "phone": emp.phone or "",
        "company_phone": (wi.mobile if wi else "") or "",
        "company_email": (wi.email if wi else "") or "",
        "personal_email": emp.email or "",
        "address": emp.address or "",
        "temporary_address": (prof.temporary_address if prof else "") or "",
        "emergency_contact_name": emp.emergency_contact_name or "",
        "emergency_contact": emp.emergency_contact or "",
        # Tuyển dụng / Onboarding
        "recruit_applied_date": hm.get("recruit_applied_date") or "",
        "recruit_source": hm.get("recruit_source") or "",
        "date_joining": _d(wi.date_joining) if wi else None,
        "probation_end": hm.get("probation_end") or "",
        "probation_salary": hm.get("probation_salary") or "",
        "reporting_manager": _vn_name(rm) if rm else "",
        "reporting_manager_id": rm.id if rm else None,
        "probation_status": hm.get("probation_status") or "",
        # Hợp đồng & công việc
        "job_position": pos.job_position if pos else "",
        "job_position_id": pos.id if pos else None,
        "level_label": hm.get("level_label") or (emp.work_level.name if getattr(emp, "work_level", None) else ""),
        "location": (wi.location if wi else "") or "",
        "contract1_type": hm.get("contract1_type") or "",
        "contract1_no": hm.get("contract1_no") or "",
        "contract1_start": hm.get("contract1_start") or "",
        "contract1_end": hm.get("contract1_end") or "",
        "contract2_type": hm.get("contract2_type") or "",
        "contract2_no": hm.get("contract2_no") or "",
        "contract2_start": hm.get("contract2_start") or "",
        "contract2_end": hm.get("contract2_end") or "",
        "contract3_type": hm.get("contract3_type") or "",
        "contract3_no": hm.get("contract3_no") or "",
        "contract3_start": hm.get("contract3_start") or "",
        "work_form": hm.get("work_form") or "",
        # Lương & phúc lợi
        "gross_salary": (wi.basic_salary if wi else None),
        "allowance": hm.get("allowance") or "",
        "pay_method": hm.get("pay_method") or "",
        "bank_account": (bank.account_number if bank else "") or "",
        "bank_name": (bank.bank_name if bank else "") or "",
        "tax_code": (prof.tax_code if prof else "") or "",
        "dependents_count": emp.children if emp.children is not None else "",
        "dependents_names": hm.get("dependents_names") or "",
        "bhxh_number": (prof.bhxh_number if prof else "") or "",
        # Học vấn
        "education": emp.qualification or "",
        "major": (prof.major if prof else "") or "",
        "school": hm.get("school") or "",
        # Đánh giá & phát triển
        "eval_rating": hm.get("eval_rating") or "",
        "raise_date": hm.get("raise_date") or "",
        "raise_amount": hm.get("raise_amount") or "",
        "raise_effective": hm.get("raise_effective") or "",
        "promotion_date": hm.get("promotion_date") or "",
        # Offboarding
        "work_status": hm.get("work_status") or ("Đang làm việc" if emp.is_active else "Đã nghỉ việc"),
        "is_active": emp.is_active,
        "resign_date": hm.get("resign_date") or "",
        "resign_reason": hm.get("resign_reason") or "",
        "resign_type": hm.get("resign_type") or "",
        "notice_date": hm.get("notice_date") or "",
        "handover": hm.get("handover") or "",
        "note": hm.get("note") or "",
        "deactivated_date": hm.get("deactivated_date") or "",
        "employee_type": (wi.employee_type_id.employee_type if wi and wi.employee_type_id else ""),
        # Đang thử việc = còn active + vào làm trong vòng PROBATION_DAYS ngày (theo
        # thông tin NV, không theo HĐ). Dùng để đánh dấu/lọc đồng nhất.
        "is_probation": _is_probation(emp, wi),
    }


def _is_probation(emp, wi) -> bool:
    from datetime import date
    dj = wi.date_joining if wi else None
    return bool(emp.is_active and dj and (date.today() - dj).days <= PROBATION_DAYS)


class HRMasterPagination(PageNumberPagination):
    page_size = 50
    page_size_query_param = "page_size"
    max_page_size = 500


def _filtered_qs(request):
    qs = Employee.objects.select_related(
        "employee_work_info__department_id",
        "employee_work_info__company_id",
        "employee_work_info__job_position_id",
        "employee_work_info__reporting_manager_id",
        "hnh_profile",
        "employee_bank_details",
        "work_level",
    )
    active = request.query_params.get("active")
    if active == "1":
        qs = qs.filter(is_active=True)
    elif active == "0":
        qs = qs.filter(is_active=False)
    # Lọc "Đang thử việc" theo THÔNG TIN NHÂN VIÊN: ngày vào làm trong vòng
    # PROBATION_DAYS gần đây. KHÔNG dùng loại HĐ (nhiều NV vào 2025 vẫn để HĐ Thử
    # việc do không cập nhật → sai); cấp NV không có field probation nào khác.
    if request.query_params.get("probation") == "1":
        from datetime import date, timedelta
        cutoff = date.today() - timedelta(days=PROBATION_DAYS)
        qs = qs.filter(
            is_active=True,
            employee_work_info__date_joining__isnull=False,
            employee_work_info__date_joining__gte=cutoff,
        )
    q = (request.query_params.get("q") or "").strip()
    if q:
        qs = qs.filter(
            Q(employee_first_name__icontains=q)
            | Q(employee_last_name__icontains=q)
            | Q(badge_id__icontains=q)
            | Q(employee_code__icontains=q)
            | Q(accounting_code__icontains=q)
        )
    company_id = request.query_params.get("company_id")
    if company_id:
        qs = qs.filter(employee_work_info__company_id=company_id)
    dept_id = request.query_params.get("department_id")
    if dept_id:
        qs = qs.filter(employee_work_info__department_id=dept_id)
    year = request.query_params.get("join_year")
    if year:
        qs = qs.filter(employee_work_info__date_joining__year=year)
    return qs.order_by("employee_work_info__company_id", "badge_id", "id")


class HRMasterDataListView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not _can_view(request):
            return Response({"detail": "Không có quyền"}, status=403)
        qs = _filtered_qs(request)
        paginator = HRMasterPagination()
        page = paginator.paginate_queryset(qs, request)
        rows = [serialize_employee(e) for e in page]
        resp = paginator.get_paginated_response(rows)
        # Kèm meta lọc: công ty, phòng ban, các năm vào công ty (cho dropdown).
        from base.models import Company, Department
        resp.data["companies"] = [{"id": c.id, "name": c.company} for c in Company.objects.all()]
        resp.data["departments"] = [
            {"id": d.id, "name": d.department, "company_ids": [c.id for c in d.company_id.all()]}
            for d in Department.objects.prefetch_related("company_id").all()
        ]
        years = sorted({
            dj.year for dj in EmployeeWorkInformation.objects.filter(
                date_joining__isnull=False
            ).values_list("date_joining", flat=True)
        }, reverse=True)
        resp.data["join_years"] = years
        resp.data["can_edit"] = _can_edit(request)
        return resp


class HRMasterDataDetailView(APIView):
    permission_classes = [IsAuthenticated]

    def _emp(self, pk):
        return Employee.objects.select_related(
            "employee_work_info__department_id", "employee_work_info__company_id",
            "employee_work_info__job_position_id", "employee_work_info__reporting_manager_id",
            "hnh_profile", "employee_bank_details", "work_level",
        ).filter(id=pk).first()

    def get(self, request, pk):
        if not _can_view(request):
            return Response({"detail": "Không có quyền"}, status=403)
        emp = self._emp(pk)
        if emp is None:
            return Response({"detail": "Không tìm thấy"}, status=404)
        data = serialize_employee(emp)
        data["can_edit"] = _can_edit(request)
        return Response(data)

    def put(self, request, pk):
        if not _can_edit(request):
            return Response({"detail": "Không có quyền sửa"}, status=403)
        emp = self._emp(pk)
        if emp is None:
            return Response({"detail": "Không tìm thấy"}, status=404)
        d = request.data

        def g(key):
            return d.get(key) if key in d else _UNSET

        try:
            with transaction.atomic():
                # ── Employee ──
                emp_fields = {
                    "employee_first_name": "first_name", "employee_last_name": "last_name",
                    "dob": "dob", "gender": "gender", "marital_status": "marital_status",
                    "phone": "phone", "email": "personal_email", "address": "address",
                    "emergency_contact_name": "emergency_contact_name",
                    "emergency_contact": "emergency_contact", "qualification": "education",
                    "accounting_code": "accounting_code", "country": "nationality",
                }
                emp_dirty = False
                for model_f, in_f in emp_fields.items():
                    if in_f in d:
                        val = d.get(in_f)
                        setattr(emp, model_f, val if val != "" else None if model_f in ("dob",) else val)
                        emp_dirty = True
                if "dependents_count" in d:
                    try:
                        emp.children = int(d.get("dependents_count")) if str(d.get("dependents_count")).strip() != "" else None
                    except (TypeError, ValueError):
                        emp.children = None
                    emp_dirty = True
                # Chỉ đổi Mã NV khi THỰC SỰ khác giá trị hiện tại (tránh ghi thừa
                # + tự va constraint). Kiểm tra trùng với NV khác → báo lỗi rõ.
                if "badge_id" in d and d.get("badge_id"):
                    new_badge = str(d.get("badge_id")).strip()
                    if new_badge and new_badge != (emp.badge_id or ""):
                        dup = Employee.objects.filter(badge_id=new_badge).exclude(id=emp.id).first()
                        if dup is not None:
                            raise _DupBadge(
                                f"Mã NV '{new_badge}' đã được dùng cho nhân viên khác "
                                f"({_vn_name(dup)} — id {dup.id}). Vui lòng dùng mã khác."
                            )
                        emp.badge_id = new_badge
                        emp.employee_code = new_badge
                        emp_dirty = True
                # extras vào additional_info['hr_master']
                ai = emp.additional_info if isinstance(emp.additional_info, dict) else {}
                hm = ai.get("hr_master") if isinstance(ai.get("hr_master"), dict) else {}
                for f in EXTRA_FIELDS:
                    if f in d:
                        hm[f] = d.get(f)
                        emp_dirty = True
                if emp_dirty:
                    ai["hr_master"] = hm
                    emp.additional_info = ai
                    emp.save()

                # ── WorkInformation ──
                wi = getattr(emp, "employee_work_info", None) or EmployeeWorkInformation(employee_id=emp)
                wi_dirty = False
                if "company_email" in d:
                    wi.email = d.get("company_email") or None; wi_dirty = True
                if "company_phone" in d:
                    wi.mobile = d.get("company_phone") or None; wi_dirty = True
                if "location" in d:
                    wi.location = d.get("location") or None; wi_dirty = True
                if "date_joining" in d:
                    wi.date_joining = d.get("date_joining") or None; wi_dirty = True
                if "gross_salary" in d:
                    try:
                        wi.basic_salary = int(d.get("gross_salary")) if str(d.get("gross_salary")).strip() != "" else 0
                    except (TypeError, ValueError):
                        wi.basic_salary = 0
                    wi_dirty = True
                if "company_id" in d and d.get("company_id"):
                    from base.models import Company
                    c = Company.objects.filter(id=d.get("company_id")).first()
                    if c: wi.company_id = c; wi_dirty = True
                if "department_id" in d and d.get("department_id"):
                    from base.models import Department
                    dep = Department.objects.filter(id=d.get("department_id")).first()
                    if dep: wi.department_id = dep; wi_dirty = True
                if "job_position_id" in d and d.get("job_position_id"):
                    from base.models import JobPosition
                    jp = JobPosition.objects.filter(id=d.get("job_position_id")).first()
                    if jp: wi.job_position_id = jp; wi_dirty = True
                if "reporting_manager_id" in d:
                    rid = d.get("reporting_manager_id")
                    wi.reporting_manager_id = Employee.objects.filter(id=rid).first() if rid else None
                    wi_dirty = True
                if wi_dirty:
                    wi.save()

                # ── HNHEmployeeProfile ──
                prof, _ = HNHEmployeeProfile.objects.get_or_create(employee_id=emp)
                prof_map = {
                    "cccd": "cccd", "cccd_issue_date": "cccd_issue_date",
                    "cccd_issue_place": "cccd_issue_place", "temporary_address": "temporary_address",
                    "tax_code": "tax_code", "bhxh_number": "bhxh_number", "major": "major",
                }
                prof_dirty = False
                for model_f, in_f in prof_map.items():
                    if in_f in d:
                        setattr(prof, model_f, d.get(in_f) or None)
                        prof_dirty = True
                if prof_dirty:
                    prof.save()

                # ── Bank ──
                if "bank_name" in d or "bank_account" in d:
                    bank = getattr(emp, "employee_bank_details", None) or EmployeeBankDetails(employee_id=emp)
                    if "bank_name" in d:
                        bank.bank_name = d.get("bank_name") or ""
                    if "bank_account" in d:
                        bank.account_number = d.get("bank_account") or ""
                    bank.save()
        except _DupBadge as e:
            return Response({"detail": str(e)}, status=400)
        except Exception as e:  # noqa: BLE001
            msg = str(e)
            if "unique_badge_id" in msg or "badge_id" in msg.lower() and "unique" in msg.lower():
                msg = "Mã NV này đã được dùng cho nhân viên khác. Vui lòng dùng mã khác."
            return Response({"detail": f"Lỗi khi lưu: {msg}"}, status=400)

        emp = self._emp(pk)
        out = serialize_employee(emp)
        out["can_edit"] = True
        return Response(out)


_UNSET = object()


class _DupBadge(Exception):
    """Mã NV bị trùng với nhân viên khác."""


# Thứ tự + nhãn cột cho Excel (khớp 69 cột Master Data).
EXPORT_COLS = [
    ("STT", None), ("company_code", "Mã Cty"), ("company", "Tên Công ty"),
    ("dept_code", "Mã Phòng ban"), ("department", "Phòng Ban"), ("team", "Bộ phận/Team"),
    ("badge_id", "Mã NV"), ("name", "Họ và tên"), ("dob", "Ngày sinh"), ("gender", "Giới tính"),
    ("cccd", "Số CCCD"), ("cccd_issue_date", "Ngày cấp CCCD"), ("cccd_issue_place", "Nơi cấp CCCD"),
    ("nationality", "Quốc tịch"), ("marital_status", "Tình trạng hôn nhân"),
    ("phone", "Số ĐT cá nhân"), ("company_phone", "Số ĐT Cty"), ("company_email", "Email công ty"),
    ("personal_email", "Email cá nhân"), ("address", "Địa chỉ thường trú"),
    ("temporary_address", "Địa chỉ hiện tại"), ("emergency_contact_name", "Người liên hệ khẩn cấp"),
    ("emergency_contact", "SĐT khẩn cấp"), ("recruit_applied_date", "Ngày ứng tuyển"),
    ("recruit_source", "Nguồn tuyển dụng"), ("date_joining", "Ngày nhận việc"),
    ("probation_end", "Ngày kết thúc thử việc"), ("probation_salary", "Lương thử việc"),
    ("reporting_manager", "Quản lý trực tiếp"), ("probation_status", "Trạng thái thử việc"),
    ("job_position", "Chức danh"), ("level_label", "Cấp bậc"), ("location", "Nơi làm việc"),
    ("contract1_type", "Loại HĐLĐ (Lần 1)"), ("contract1_no", "Số HĐ (Lần 1)"),
    ("contract1_start", "Ngày bắt đầu HĐ 1"), ("contract1_end", "Ngày hết hạn HĐ 1"),
    ("contract2_type", "Loại HĐLĐ (Lần 2)"), ("contract2_no", "Số HĐ (Lần 2)"),
    ("contract2_start", "Ngày bắt đầu HĐ 2"), ("contract2_end", "Ngày hết hạn HĐ 2"),
    ("contract3_type", "Loại HĐLĐ (Lần 3)"), ("contract3_no", "Số HĐ (Lần 3)"),
    ("contract3_start", "Ngày bắt đầu HĐ 3"), ("work_form", "Hình thức làm việc"),
    ("gross_salary", "Lương Gross (VND)"), ("allowance", "Phụ cấp"), ("pay_method", "Hình thức trả lương"),
    ("bank_account", "Số tài khoản"), ("bank_name", "Ngân hàng"), ("tax_code", "Mã số thuế TNCN"),
    ("dependents_count", "SL người phụ thuộc"), ("dependents_names", "Tên người phụ thuộc"),
    ("bhxh_number", "Số sổ BHXH"), ("education", "Trình độ học vấn"), ("major", "Chuyên ngành"),
    ("school", "Trường tốt nghiệp"), ("eval_rating", "Xếp loại đánh giá gần nhất"),
    ("raise_date", "Ngày tăng lương gần nhất"), ("raise_amount", "Mức tăng"),
    ("raise_effective", "Ngày hiệu lực"), ("promotion_date", "Ngày thăng chức gần nhất"),
    ("work_status", "Trạng thái làm việc"), ("resign_date", "Ngày nghỉ việc"),
    ("resign_reason", "Lý do nghỉ việc"), ("resign_type", "Loại nghỉ việc"),
    ("notice_date", "Ngày báo trước nghỉ việc"), ("handover", "Bàn giao công việc"), ("note", "Ghi chú"),
]


class HRMasterDataExportView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not _can_view(request):
            return Response({"detail": "Không có quyền"}, status=403)
        from django.http import HttpResponse
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        import io

        rows = [serialize_employee(e) for e in _filtered_qs(request)[:5000]]
        wb = Workbook()
        ws = wb.active
        ws.title = "Master Data Nhân sự"
        thin = Border(*(Side(style="thin", color="D9D9D9"),) * 4)
        hdr_font = Font(bold=True, color="FFFFFF", size=9)
        hdr_fill = PatternFill(start_color="1A2340", end_color="1A2340", fill_type="solid")
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col, (key, label) in enumerate(EXPORT_COLS, 1):
            c = ws.cell(row=1, column=col, value=label if label else "STT")
            c.font = hdr_font; c.fill = hdr_fill; c.alignment = center; c.border = thin
            ws.column_dimensions[c.column_letter].width = 8 if key in (None, "STT") else 16
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "I2"  # cố định STT→Họ tên + hàng tiêu đề
        for i, row in enumerate(rows, 1):
            ws.cell(row=i + 1, column=1, value=i).border = thin
            for col, (key, _label) in enumerate(EXPORT_COLS[1:], 2):
                v = row.get(key, "")
                c = ws.cell(row=i + 1, column=col, value="" if v is None else v)
                c.border = thin
                c.alignment = Alignment(vertical="center")
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        resp = HttpResponse(
            buf.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp["Content-Disposition"] = 'attachment; filename="MasterData_NhanSu.xlsx"'
        return resp


def _parse_any_date(s):
    """Parse resign_date/deactivated_date (ISO 'YYYY-MM-DD', 'DD/MM/YYYY', ISO datetime)."""
    if not s:
        return None
    from datetime import datetime as _dt
    s = str(s).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return _dt.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    try:
        return _dt.fromisoformat(s).date()
    except ValueError:
        return None


class HRMasterLeaversExportView(APIView):
    """Xuất Excel NV tạm dừng/nghỉ việc TRONG THÁNG.

    GET /api/employee/hr-master/export-leavers/?month=YYYY-MM (mặc định tháng hiện tại).
    Ngày hiệu lực = resign_date (HR nhập, ưu tiên) → else deactivated_date (hệ thống
    tự ghi khi C&B chuyển Tạm nghỉ). Chỉ NV is_active=False có ngày rơi trong tháng."""

    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not _can_view(request):
            return Response({"detail": "Không có quyền"}, status=403)
        from django.http import HttpResponse
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from datetime import date
        import io, calendar

        try:
            y, m = map(int, (request.query_params.get("month") or "").split("-"))
            assert 1 <= m <= 12
        except Exception:
            t = date.today(); y, m = t.year, t.month
        start, end = date(y, m, 1), date(y, m, calendar.monthrange(y, m)[1])

        rows = []
        for e in Employee.objects.filter(is_active=False).select_related(
            "employee_work_info__department_id", "employee_work_info__company_id",
            "employee_work_info__job_position_id", "employee_work_info__employee_type_id",
        ):
            hm = _extra(e)
            eff = _parse_any_date(hm.get("resign_date")) or _parse_any_date(hm.get("deactivated_date"))
            if not eff or not (start <= eff <= end):
                continue
            wi = getattr(e, "employee_work_info", None)
            rows.append({
                "badge": e.badge_id or e.employee_code or "", "name": _vn_name(e),
                "company": wi.company_id.company if wi and wi.company_id else "",
                "dept": wi.department_id.department if wi and wi.department_id else "",
                "pos": wi.job_position_id.job_position if wi and wi.job_position_id else "",
                "etype": wi.employee_type_id.employee_type if wi and wi.employee_type_id else "",
                "kind": "Nghỉ việc" if hm.get("resign_date") else "Tạm dừng",
                "eff": eff.strftime("%d/%m/%Y"),
                "rtype": hm.get("resign_type") or "", "reason": hm.get("resign_reason") or "",
            })
        rows.sort(key=lambda r: r["eff"])

        wb = Workbook(); ws = wb.active; ws.title = f"NV nghỉ {m:02d}-{y}"
        cols = [("STT", 6), ("Mã NV", 12), ("Họ và tên", 24), ("Công ty", 22),
                ("Phòng ban", 20), ("Chức danh", 20), ("Loại HĐ", 14), ("Loại nghỉ", 11),
                ("Ngày nghỉ", 12), ("Hình thức", 16), ("Lý do", 30)]
        thin = Border(*(Side(style="thin", color="D9D9D9"),) * 4)
        hf = Font(bold=True, color="FFFFFF", size=10)
        hfill = PatternFill("solid", fgColor="1A2340")
        ce = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
        t = ws.cell(row=1, column=1, value=f"DANH SÁCH NHÂN VIÊN TẠM DỪNG / NGHỈ VIỆC — THÁNG {m:02d}/{y}")
        t.font = Font(bold=True, size=13, color="C0222B"); t.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24
        for col, (name, w) in enumerate(cols, 1):
            c = ws.cell(row=2, column=col, value=name)
            c.font = hf; c.fill = hfill; c.alignment = ce; c.border = thin
            ws.column_dimensions[c.column_letter].width = w
        for i, r in enumerate(rows, 1):
            vals = [i, r["badge"], r["name"], r["company"], r["dept"], r["pos"],
                    r["etype"], r["kind"], r["eff"], r["rtype"], r["reason"]]
            for col, v in enumerate(vals, 1):
                c = ws.cell(row=i + 2, column=col, value=v)
                c.border = thin; c.alignment = Alignment(vertical="center", wrap_text=(col in (3, 11)))
        ws.freeze_panes = "A3"
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        resp = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = f'attachment; filename="NV_nghi_{y}_{m:02d}.xlsx"'
        return resp


class HRCategoriesView(APIView):
    """Danh mục: công ty, phòng ban (DB) + vị trí + 13 danh mục enum."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not _can_view(request):
            return Response({"detail": "Không có quyền"}, status=403)
        from base.models import Company, Department, JobPosition

        companies = [{"id": c.id, "name": c.company, "address": c.address or ""} for c in Company.objects.all()]
        departments = [
            {"id": d.id, "name": d.department, "company_ids": [c.id for c in d.company_id.all()]}
            for d in Department.objects.prefetch_related("company_id").all()
        ]
        positions = [
            {"id": p.id, "name": p.job_position,
             "department": p.department_id.department if p.department_id else ""}
            for p in JobPosition.objects.select_related("department_id").all()
        ]
        enums = [
            {"key": k, "label": ENUM_LABELS.get(k, k), "values": v}
            for k, v in ENUM_CATEGORIES.items()
        ]
        return Response({
            "companies": companies,
            "departments": departments,
            "positions": positions,
            "enums": enums,
        })
